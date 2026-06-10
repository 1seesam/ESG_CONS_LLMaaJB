# -*- coding: utf-8 -*-
from __future__ import annotations

import os
import json
import subprocess
import sys
from typing import Optional

import docx
import extract_msg
import pandas as pd
from langchain_community.document_loaders import PyPDFLoader
from pptx import Presentation

# Default input directory
DATA_FOLDER = "./RAG_data"
DOCLING_TIMEOUT_SEC = 60


def _docling_available() -> bool:
    try:
        import docling  # type: ignore  # noqa: F401

        return True
    except Exception:
        return False


def _read_pdf_docling(filepath: str) -> str:
    from docling.document_converter import DocumentConverter  # type: ignore

    converter = DocumentConverter()
    result = converter.convert(filepath)
    doc = getattr(result, "document", result)

    for method_name in ("export_to_markdown", "export_to_text", "to_markdown"):
        method = getattr(doc, method_name, None)
        if callable(method):
            value = str(method()).strip()
            if value:
                return value

    text_attr = getattr(doc, "text", None)
    if isinstance(text_attr, str) and text_attr.strip():
        return text_attr.strip()
    return str(doc).strip()


def _read_pdf_docling_timeout(filepath: str, timeout_sec: int = DOCLING_TIMEOUT_SEC) -> str:
    script = (
        "import json,sys\n"
        "from docling.document_converter import DocumentConverter\n"
        "path=sys.argv[1]\n"
        "result=DocumentConverter().convert(path)\n"
        "doc=getattr(result,'document',result)\n"
        "text=''\n"
        "for name in ('export_to_markdown','export_to_text','to_markdown'):\n"
        "    fn=getattr(doc,name,None)\n"
        "    if callable(fn):\n"
        "        value=str(fn()).strip()\n"
        "        if value:\n"
        "            text=value\n"
        "            break\n"
        "if not text:\n"
        "    t=getattr(doc,'text',None)\n"
        "    text=t.strip() if isinstance(t,str) else str(doc).strip()\n"
        "print(json.dumps({'text': text}, ensure_ascii=False))\n"
    )
    completed = subprocess.run(
        [sys.executable, "-c", script, filepath],
        capture_output=True,
        text=True,
        encoding="utf-8",
        timeout=max(1, int(timeout_sec)),
    )
    if completed.returncode != 0:
        err = completed.stderr.strip() or completed.stdout.strip() or "docling subprocess failed"
        raise RuntimeError(err)
    payload = json.loads(completed.stdout.strip() or "{}")
    return str(payload.get("text", "")).strip()


def _read_pdf_pypdf(filepath: str) -> str:
    try:
        loader = PyPDFLoader(filepath)
        pages = loader.load()
        text = ""
        for i, page in enumerate(pages):
            text += f"\n--- PDF SEITE {i+1} ---\n{page.page_content}\n"
        return text
    except Exception as exc:
        return f"[PDF FEHLER: {exc}]"


def read_pdf(filepath: str, pdf_reader: Optional[str] = "auto") -> str:
    """Read PDF using Docling when available, otherwise PyPDFLoader fallback."""
    pref = str(pdf_reader or "auto").strip().lower()
    use_docling = _docling_available() if pref == "auto" else pref == "docling"

    if use_docling:
        try:
            text = _read_pdf_docling_timeout(filepath, timeout_sec=DOCLING_TIMEOUT_SEC)
            if text:
                return text
        except Exception:
            if pref == "docling":
                return _read_pdf_pypdf(filepath)

    return _read_pdf_pypdf(filepath)


def read_excel(filepath: str) -> str:
    try:
        xls = pd.ExcelFile(filepath)
        text = ""
        for sheet_name in xls.sheet_names:
            df = pd.read_excel(xls, sheet_name=sheet_name)
            df = df.dropna(how="all").dropna(axis=1, how="all")
            tbl_text = df.to_markdown(index=False)
            text += f"\n--- EXCEL BLATT: '{sheet_name}' ---\n{tbl_text}\n"
        return text
    except Exception as exc:
        return f"[EXCEL FEHLER: {exc}]"


def read_word(filepath: str) -> str:
    try:
        doc = docx.Document(filepath)
        full_text = []
        full_text.append("--- FLIESSTEXT ---")
        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text)

        full_text.append("\n--- TABELLEN IN DOKUMENT ---")
        for table in doc.tables:
            for row in table.rows:
                row_data = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                full_text.append(" | ".join(row_data))
            full_text.append("[Tabelle Ende]\n")
        return "\n".join(full_text)
    except Exception as exc:
        return f"[WORD FEHLER: {exc}]"


def read_powerpoint(filepath: str) -> str:
    try:
        prs = Presentation(filepath)
        text = ""
        for i, slide in enumerate(prs.slides):
            text += f"\n--- FOLIE {i+1} ---\n"
            if slide.shapes.title:
                text += f"TITEL: {slide.shapes.title.text}\n"
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    if shape == slide.shapes.title:
                        continue
                    text += f"- {shape.text.replace(chr(11), ' ')}\n"
        return text
    except Exception as exc:
        return f"[PPTX FEHLER: {exc}]"


def read_msg(filepath: str) -> str:
    try:
        msg = extract_msg.Message(filepath)
        text = f"Von: {msg.sender}\nAn: {msg.to}\nBetreff: {msg.subject}\nDatum: {msg.date}\n"
        text += f"\n--- MAIL BODY ---\n{msg.body}\n"
        msg.close()
        return text
    except Exception as exc:
        return f"[MAIL FEHLER: {exc}]"


def get_full_document_text(filename: str, data_folder: str = DATA_FOLDER) -> str:
    """Read a file based on extension and return normalized text blob."""
    if not filename or str(filename).lower() == "nan":
        return ""

    filepath = os.path.join(data_folder, filename)

    if not os.path.exists(filepath):
        found = False
        for ext in [".pdf", ".xlsx", ".xls", ".docx", ".pptx", ".msg", ".txt"]:
            if os.path.exists(filepath + ext):
                filepath += ext
                filename += ext
                found = True
                break
        if not found:
            return f"[FEHLER: Datei '{filename}' nicht in '{data_folder}' gefunden]"

    _, ext = os.path.splitext(filename)
    ext = ext.lower()

    header = f"=== DOKUMENT START: {filename} ===\n"
    content = ""

    if ext == ".pdf":
        content = read_pdf(filepath, pdf_reader="auto")
    elif ext in [".xlsx", ".xls"]:
        content = read_excel(filepath)
    elif ext == ".docx":
        content = read_word(filepath)
    elif ext == ".pptx":
        content = read_powerpoint(filepath)
    elif ext == ".msg":
        content = read_msg(filepath)
    elif ext == ".txt":
        with open(filepath, "r", encoding="utf-8") as f:
            content = f.read()
    else:
        content = f"[FORMAT NICHT UNTERSTUETZT: {ext}]"

    return header + content + "\n=== DOKUMENT ENDE ===\n"
